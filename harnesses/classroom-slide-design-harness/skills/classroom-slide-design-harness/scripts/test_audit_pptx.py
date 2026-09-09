#!/usr/bin/env python3
from __future__ import annotations

import json
import io
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parent))
import audit_pptx  # noqa: E402


class AuditPptxTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)

    def blank_deck(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        return presentation, slide

    def test_text_and_table_cells_are_counted(self):
        presentation, slide = self.blank_deck()
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = "Visible text here"
        table = slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(6), Inches(1)).table
        table.cell(0, 0).text, table.cell(0, 1).text = "table one", "table two words"
        presentation.save(self.root / "deck.pptx")
        result = audit_pptx.audit_pptx(self.root / "deck.pptx", self.root / "output")
        self.assertEqual(result["slide_count"], 1)
        self.assertEqual(result["visible_word_count"], 8)
        self.assertEqual(result["slides"][0]["table_cell_count"], 2)
        saved = json.loads((self.root / "output" / "audit.json").read_text(encoding="utf-8"))
        self.assertEqual(saved, result)

    def test_nested_group_text_is_counted(self):
        presentation, slide = self.blank_deck()
        outer = slide.shapes.add_group_shape()
        inner = outer.shapes.add_group_shape()
        inner.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = "nested group words"
        presentation.save(self.root / "deck.pptx")
        result = audit_pptx.audit_pptx(self.root / "deck.pptx", self.root / "output")
        self.assertEqual(result["visible_word_count"], 3)
        self.assertGreaterEqual(result["slides"][0]["shape_count"], 3)

    def test_out_of_slide_geometry_fails(self):
        presentation, slide = self.blank_deck()
        slide.shapes.add_shape(MSO_SHAPE.OVAL, presentation.slide_width - Inches(0.25),
                               Inches(1), Inches(1), Inches(1))
        presentation.save(self.root / "deck.pptx")
        result = audit_pptx.audit_pptx(self.root / "deck.pptx", self.root / "output")
        self.assertEqual(result["verdict"], "FAIL")
        self.assertEqual(result["structural_issues"][0]["check"], "out_of_slide_geometry")

    def test_input_hash_is_unchanged(self):
        presentation, _ = self.blank_deck()
        presentation.save(self.root / "deck.pptx")
        before = audit_pptx.sha256_file(self.root / "deck.pptx")
        result = audit_pptx.audit_pptx(self.root / "deck.pptx", self.root / "output")
        self.assertEqual(audit_pptx.sha256_file(self.root / "deck.pptx"), before)
        self.assertTrue(result["source"]["unchanged"])

    def test_reused_output_and_input_collision_are_rejected(self):
        presentation, _ = self.blank_deck()
        presentation.save(self.root / "deck.pptx")
        audit_pptx.audit_pptx(self.root / "deck.pptx", self.root / "output")
        with self.assertRaises(audit_pptx.AuditError):
            audit_pptx.audit_pptx(self.root / "deck.pptx", self.root / "output")
        with self.assertRaises(audit_pptx.AuditError):
            audit_pptx.audit_pptx(self.root / "deck.pptx", self.root / "deck.pptx", overwrite=True)
        with self.assertRaises(audit_pptx.AuditError):
            audit_pptx.audit_pptx(self.root / "deck.pptx", self.root, overwrite=True)

    def test_overwrite_preserves_unrelated_files(self):
        presentation, _ = self.blank_deck()
        source, output = self.root / "deck.pptx", self.root / "output"
        presentation.save(source)
        audit_pptx.audit_pptx(source, output)
        (output / "student-notes.txt").write_text("keep", encoding="utf-8")
        (output / "office-slide-999.png").write_bytes(b"stale generated image")
        result = audit_pptx.audit_pptx(source, output, overwrite=True)
        self.assertEqual(result["verdict"], "PASS")
        self.assertEqual((output / "student-notes.txt").read_text(encoding="utf-8"), "keep")
        self.assertFalse((output / "office-slide-999.png").exists())

    def test_cli_returns_failure_for_out_of_slide_content(self):
        presentation, slide = self.blank_deck()
        slide.shapes.add_textbox(-Inches(1), Inches(1), Inches(2), Inches(1)).text = "outside"
        source = self.root / "deck.pptx"
        presentation.save(source)
        with mock.patch("sys.stdout", new_callable=io.StringIO):
            status = audit_pptx.main([str(source), "--output-dir", str(self.root / "output")])
        self.assertEqual(status, 1)

    def test_requested_office_render_unavailable_is_a_failure(self):
        presentation, _ = self.blank_deck()
        presentation.save(self.root / "deck.pptx")
        with mock.patch.object(audit_pptx.sys, "platform", "portable-test"):
            result = audit_pptx.audit_pptx(
                self.root / "deck.pptx", self.root / "output", render_office=True)
        self.assertEqual(result["office_render"]["status"], "unavailable")
        self.assertEqual(result["verdict"], "FAIL")
        self.assertTrue((self.root / "output" / "audit.json").is_file())

    def test_office_measurement_detects_text_overflow(self):
        shape = mock.Mock()
        shape.HasTextFrame = shape.TextFrame.HasText = True
        shape.TextFrame.TextRange.Text = "A visible label"
        shape.Width, shape.Height = 120, 36
        shape.TextFrame.MarginLeft = shape.TextFrame.MarginRight = 0
        shape.TextFrame.MarginTop = shape.TextFrame.MarginBottom = 0
        shape.TextFrame.TextRange.BoundWidth = 100
        shape.TextFrame.TextRange.BoundHeight = 40
        findings = []
        audit_pptx._measure_text_frame(shape, 1, "label", findings)
        self.assertEqual(findings[0]["dimensions"], ["height"])
        shape.Height = 50
        findings = []
        audit_pptx._measure_text_frame(shape, 1, "label", findings)
        self.assertEqual(findings, [])

    def test_office_cleanup_preserves_other_open_presentations(self):
        presentation, _ = self.blank_deck()
        presentation.save(self.root / "deck.pptx")
        opened, application = mock.Mock(), mock.Mock()
        opened.Slides.Count = 0
        application.Presentations.Open.return_value = opened
        application.Presentations.Count = 1
        dispatch_ex = mock.Mock(return_value=application)
        with mock.patch.object(audit_pptx.sys, "platform", "win32"), mock.patch.object(audit_pptx, "_make_contact_sheet"):
            result = audit_pptx._render_with_office(self.root / "deck.pptx", self.root, dispatch_ex)
        application.Presentations.Open.assert_called_once_with(
            FileName=str(self.root / "deck.pptx"), ReadOnly=True, Untitled=False, WithWindow=False)
        opened.Close.assert_called_once_with()
        opened.SaveAs.assert_called_once_with(str(self.root / "office-render.pdf"), 32)
        application.Quit.assert_not_called()
        self.assertEqual(result["application_quit"], "skipped_presentations_open")


if __name__ == "__main__":
    unittest.main()
