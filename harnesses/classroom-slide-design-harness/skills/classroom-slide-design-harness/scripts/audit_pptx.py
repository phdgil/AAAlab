#!/usr/bin/env python3
"""Read-only structural PPTX audit with optional Microsoft PowerPoint rendering."""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import sys
from pathlib import Path
from typing import Any, Iterable

from PIL import Image, ImageDraw, ImageOps
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

WORD_RE = re.compile(r"\b\w+(?:[’'-]\w+)*\b", re.UNICODE)
EMU_PER_POINT = 12700
OWNED_OUTPUTS = ("audit.json", "office-render.pdf", "office-contact-sheet.png")


class AuditError(RuntimeError):
    """The audit cannot be started safely."""


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _prepare_paths(deck_path: str | Path, output_dir: str | Path, overwrite: bool) -> tuple[Path, Path]:
    source, output = Path(deck_path).expanduser().resolve(), Path(output_dir).expanduser().resolve()
    if not source.is_file() or source.suffix.casefold() != ".pptx":
        raise AuditError(f"Input must be an existing .pptx file: {source}")
    if source == output:
        raise AuditError("The input deck and output directory collide")
    try:
        source.relative_to(output)
    except ValueError:
        pass
    else:
        raise AuditError("The output directory must not contain the input deck")
    if output.exists() and not output.is_dir():
        raise AuditError(f"Output path is not a directory: {output}")
    if output.exists() and any(output.iterdir()) and not overwrite:
        raise AuditError(f"Output directory is nonempty; use --overwrite: {output}")
    output.mkdir(parents=True, exist_ok=True)
    if overwrite:
        candidates = [output / name for name in OWNED_OUTPUTS]
        candidates.extend(output.glob("office-slide-*.png"))
        for candidate in candidates:
            if candidate.is_file() or candidate.is_symlink():
                candidate.unlink()
            elif candidate.exists():
                raise AuditError(f"Audit output is unexpectedly a directory: {candidate}")
    return source, output


def _child_transform(shape: Any, parent: tuple[float, float, float, float]) -> tuple[float, float, float, float]:
    ox, oy, sx, sy = parent
    try:
        xfrm = shape._element.grpSpPr.xfrm
        child_offset, child_extent = xfrm.chOff, xfrm.chExt
        if not child_extent.cx or not child_extent.cy:
            return parent
        local_sx = float(shape.width) / float(child_extent.cx)
        local_sy = float(shape.height) / float(child_extent.cy)
        return (
            ox + sx * (float(shape.left) - float(child_offset.x) * local_sx),
            oy + sy * (float(shape.top) - float(child_offset.y) * local_sy),
            sx * local_sx,
            sy * local_sy,
        )
    except (AttributeError, TypeError, ValueError, ZeroDivisionError):
        return parent


def _shape_bounds(shape: Any, transform: tuple[float, float, float, float]):
    try:
        x, y, width, height = map(float, (shape.left, shape.top, shape.width, shape.height))
    except (TypeError, ValueError):
        return None
    ox, oy, sx, sy = transform
    x1, y1 = ox + sx * x, oy + sy * y
    x2, y2 = x1 + sx * width, y1 + sy * height
    return min(x1, x2), min(y1, y2), max(x1, x2), max(y1, y2)


def _visit_shapes(
    shapes: Iterable[Any], slide_number: int, slide_size: tuple[int, int],
    transform: tuple[float, float, float, float], prefix: str, texts: list[str],
    issues: list[dict[str, Any]], metrics: dict[str, int],
) -> None:
    slide_width, slide_height = slide_size
    for index, shape in enumerate(shapes, 1):
        path = f"{prefix}/{index}:{getattr(shape, 'name', f'shape-{index}')}"
        metrics["shape_count"] += 1
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            _visit_shapes(shape.shapes, slide_number, slide_size, _child_transform(shape, transform),
                          path, texts, issues, metrics)
            continue
        bounds = _shape_bounds(shape, transform)
        if bounds is not None:
            left, top, right, bottom = bounds
            edges = [name for name, outside in (
                ("left", left < 0), ("top", top < 0),
                ("right", right > slide_width), ("bottom", bottom > slide_height),
            ) if outside]
            if edges:
                issues.append({
                    "severity": "error", "check": "out_of_slide_geometry",
                    "slide": slide_number, "shape_path": path, "outside_edges": edges,
                    "bounds_points": {key: round(value / EMU_PER_POINT, 3) for key, value in
                                      zip(("left", "top", "right", "bottom"), bounds)},
                })
        if shape.has_table:
            metrics["table_count"] += 1
            for row in shape.table.rows:
                for cell in row.cells:
                    metrics["table_cell_count"] += 1
                    texts.append(cell.text)
        elif shape.has_text_frame:
            texts.append(shape.text_frame.text)


def _measurement_error(slide: int, path: str, exc: Exception) -> dict[str, Any]:
    return {"severity": "error", "check": "office_text_measurement_error", "slide": slide,
            "shape_path": path, "error": f"{type(exc).__name__}: {exc}"}


def _measure_text_frame(shape: Any, slide: int, path: str, findings: list[dict[str, Any]]) -> None:
    try:
        if not bool(shape.HasTextFrame):
            return
        frame = shape.TextFrame
        if not bool(frame.HasText) or not str(frame.TextRange.Text).strip():
            return
        bound_w, bound_h = float(frame.TextRange.BoundWidth), float(frame.TextRange.BoundHeight)
        area_w = max(0.0, float(shape.Width) - float(frame.MarginLeft) - float(frame.MarginRight))
        area_h = max(0.0, float(shape.Height) - float(frame.MarginTop) - float(frame.MarginBottom))
        dimensions = [name for name, overflow in
                      (("width", bound_w > area_w + 0.5), ("height", bound_h > area_h + 0.5)) if overflow]
        if dimensions:
            findings.append({
                "severity": "error", "check": "office_text_overflow", "slide": slide,
                "shape_path": path, "dimensions": dimensions,
                "bound_points": {"width": round(bound_w, 3), "height": round(bound_h, 3)},
                "text_area_points": {"width": round(area_w, 3), "height": round(area_h, 3)},
            })
    except Exception as exc:
        findings.append(_measurement_error(slide, path, exc))


def _measure_office_shapes(shapes: Any, slide: int, prefix: str, findings: list[dict[str, Any]]) -> None:
    for index in range(1, int(shapes.Count) + 1):
        shape = shapes.Item(index)
        try:
            name = str(shape.Name)
        except Exception:
            name = f"shape-{index}"
        path = f"{prefix}/{index}:{name}"
        try:
            if int(shape.Type) == 6:  # msoGroup
                _measure_office_shapes(shape.GroupItems, slide, path, findings)
            elif bool(shape.HasTable):
                table = shape.Table
                for row in range(1, int(table.Rows.Count) + 1):
                    for column in range(1, int(table.Columns.Count) + 1):
                        _measure_text_frame(table.Cell(row, column).Shape, slide,
                                            f"{path}/cell-{row}-{column}", findings)
            else:
                _measure_text_frame(shape, slide, path, findings)
        except Exception as exc:
            findings.append(_measurement_error(slide, path, exc))


def _make_contact_sheet(paths: list[Path], destination: Path) -> None:
    images: list[Image.Image] = []
    sheet = None
    try:
        for path in paths:
            with Image.open(path) as image:
                images.append(image.convert("RGB"))
        if not images:
            sheet = Image.new("RGB", (640, 100), "white")
            ImageDraw.Draw(sheet).text((20, 35), "No slides", fill="black")
        else:
            columns = min(4, max(1, math.ceil(math.sqrt(len(images)))))
            rows, cell_w, cell_h = math.ceil(len(images) / columns), 340, 220
            sheet = Image.new("RGB", (columns * cell_w, rows * cell_h), "white")
            draw = ImageDraw.Draw(sheet)
            for index, image in enumerate(images):
                thumb = ImageOps.contain(image, (320, 185))
                column, row = index % columns, index // columns
                sheet.paste(thumb, (column * cell_w + (cell_w - thumb.width) // 2, row * cell_h + 10))
                draw.text((column * cell_w + 10, row * cell_h + 200), f"Slide {index + 1}", fill="black")
                thumb.close()
        sheet.save(destination, "PNG")
    finally:
        if sheet is not None:
            sheet.close()
        for image in images:
            image.close()


def _render_with_office(source: Path, output: Path, dispatch_ex: Any = None) -> dict[str, Any]:
    report: dict[str, Any] = {
        "requested": True, "renderer": "Microsoft PowerPoint via pywin32", "status": "error",
        "findings": [], "artifacts": {}, "visual_review_required": True,
    }
    if sys.platform != "win32":
        report.update(status="unavailable", error="Office rendering is available only on Windows")
        return report
    if dispatch_ex is None:
        try:
            from win32com.client import DispatchEx
            dispatch_ex = DispatchEx
        except Exception as exc:
            report.update(status="unavailable", error=f"pywin32 is unavailable: {type(exc).__name__}: {exc}")
            return report

    application = presentation = None
    png_paths: list[Path] = []
    cleanup_errors: list[str] = []
    try:
        application = dispatch_ex("PowerPoint.Application")
        presentation = application.Presentations.Open(
            FileName=str(source), ReadOnly=True, Untitled=False, WithWindow=False)
        slide_count = int(presentation.Slides.Count)
        digits = max(3, len(str(slide_count)))
        for slide_number in range(1, slide_count + 1):
            slide = presentation.Slides.Item(slide_number)
            _measure_office_shapes(slide.Shapes, slide_number, f"slide-{slide_number}", report["findings"])
            path = output / f"office-slide-{slide_number:0{digits}d}.png"
            slide.Export(str(path), "PNG")
            png_paths.append(path)
        pdf_path, contact_path = output / "office-render.pdf", output / "office-contact-sheet.png"
        presentation.SaveAs(str(pdf_path), 32)  # ppSaveAsPDF; avoids optional COM PrintRange binding
        _make_contact_sheet(png_paths, contact_path)
        report["artifacts"] = {"slide_pngs": [path.name for path in png_paths],
                               "pdf": pdf_path.name, "contact_sheet": contact_path.name}
        report["status"] = "completed"
        if any(item["check"] == "office_text_measurement_error" for item in report["findings"]):
            report.update(status="error", error="One or more Office text measurements failed")
    except Exception as exc:
        report.update(status="error", error=f"{type(exc).__name__}: {exc}")
        report["artifacts"]["slide_pngs"] = [path.name for path in png_paths]
    finally:
        if presentation is not None:
            try:
                presentation.Close()
            except Exception as exc:
                cleanup_errors.append(f"presentation.Close: {type(exc).__name__}: {exc}")
        if application is not None:
            try:
                remaining = int(application.Presentations.Count)
                if remaining == 0:
                    application.Quit()
                    report["application_quit"] = "completed"
                else:
                    report["application_quit"] = "skipped_presentations_open"
                    report["remaining_presentation_count"] = remaining
            except Exception as exc:
                cleanup_errors.append(f"application lifecycle check: {type(exc).__name__}: {exc}")
    if cleanup_errors:
        report["cleanup_errors"] = cleanup_errors
        messages = ([report["error"]] if report.get("error") else []) + cleanup_errors
        report.update(status="error", error="; ".join(messages))
    return report


def audit_pptx(deck_path: str | Path, output_dir: str | Path, *, render_office: bool = False, overwrite: bool = False) -> dict[str, Any]:
    """Audit a deck, write ``audit.json``, and return the same report."""
    source, output = _prepare_paths(deck_path, output_dir, overwrite)
    before_hash = sha256_file(source)
    presentation = Presentation(str(source))
    issues: list[dict[str, Any]] = []
    slides: list[dict[str, Any]] = []
    slide_size = (presentation.slide_width, presentation.slide_height)
    for number, slide in enumerate(presentation.slides, 1):
        texts: list[str] = []
        metrics = {"shape_count": 0, "table_count": 0, "table_cell_count": 0}
        _visit_shapes(slide.shapes, number, slide_size, (0.0, 0.0, 1.0, 1.0),
                      f"slide-{number}", texts, issues, metrics)
        words = sum(len(WORD_RE.findall(text)) for text in texts)
        slides.append({"slide": number, "visible_word_count": words, **metrics})
    office = (_render_with_office(source, output) if render_office else
              {"requested": False, "status": "not_requested", "visual_review_required": True})
    after_hash = sha256_file(source)
    integrity = ([] if before_hash == after_hash else [{
        "severity": "error", "check": "source_hash_changed",
        "before": before_hash, "after": after_hash,
    }])
    failed = bool(issues or integrity or office.get("findings") or
                  (render_office and office["status"] != "completed"))
    report: dict[str, Any] = {
        "schema_version": 1,
        "source": {"path": str(source), "sha256_before": before_hash,
                   "sha256_after": after_hash, "unchanged": before_hash == after_hash},
        "slide_count": len(slides),
        "visible_word_count": sum(item["visible_word_count"] for item in slides),
        "slides": slides, "structural_issues": issues, "integrity_issues": integrity,
        "office_render": office,
        "limitations": [
            "Structural-only inspection cannot verify rendered clipping or contrast.",
            "Structural geometry uses axis-aligned approximations and does not fully model rotation or flips.",
            "Automated auditing cannot determine learning quality.",
            "Visual design and contrast still require image review, including after Office rendering.",
            "Word counts cover text frames and table cells; embedded chart, SmartArt, media, or OLE text may be inaccessible.",
        ],
        "verdict": "FAIL" if failed else "PASS",
    }
    (output / "audit.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return report


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Audit a PPTX without modifying it.")
    parser.add_argument("deck", type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--render-office", action="store_true")
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args(argv)
    try:
        report = audit_pptx(args.deck, args.output_dir, render_office=args.render_office,
                            overwrite=args.overwrite)
    except Exception as exc:
        print(f"audit error: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 2
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["verdict"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
