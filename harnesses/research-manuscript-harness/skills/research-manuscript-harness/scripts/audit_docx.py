#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
from io import BytesIO
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_LINE_SPACING
except ImportError as exc:
    raise SystemExit("python-docx is required for DOCX auditing") from exc

try:
    from PIL import Image
except ImportError as exc:
    raise SystemExit("Pillow is required for embedded-image auditing") from exc


def all_paragraphs(document: Document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def inherited_style_value(style, accessor):
    current = style
    while current is not None:
        value = accessor(current)
        if value is not None:
            return value
        current = current.base_style
    return None


def effective_line_spacing_rule(paragraph):
    direct = paragraph.paragraph_format.line_spacing_rule
    if direct is not None:
        return direct
    return inherited_style_value(
        paragraph.style,
        lambda style: style.paragraph_format.line_spacing_rule,
    )


def effective_run_font_value(run, paragraph, attribute):
    direct = getattr(run.font, attribute)
    if direct is not None:
        return direct
    character_value = inherited_style_value(
        run.style,
        lambda style: getattr(style.font, attribute),
    )
    if character_value is not None:
        return character_value
    return inherited_style_value(
        paragraph.style,
        lambda style: getattr(style.font, attribute),
    )


def font_color_marker(font):
    if font.color.rgb is not None:
        return str(font.color.rgb).upper()
    if font.color.type is not None:
        return f"TYPE:{font.color.type}"
    return None


def effective_run_color(run, paragraph):
    direct = font_color_marker(run.font)
    if direct is not None:
        return direct
    character_value = inherited_style_value(
        run.style,
        lambda style: font_color_marker(style.font),
    )
    if character_value is not None:
        return character_value
    return inherited_style_value(
        paragraph.style,
        lambda style: font_color_marker(style.font),
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="Audit manuscript DOCX structure and author formatting rules.")
    parser.add_argument("docx", type=Path)
    parser.add_argument("--expect-double-spacing", action="store_true")
    parser.add_argument("--expect-plain-black", action="store_true")
    parser.add_argument("--expect-table-count", type=int)
    parser.add_argument(
        "--expect-table-shape",
        action="append",
        default=[],
        metavar="ROWSxCOLS",
        help="Expected table shape in document order; repeat for each table.",
    )
    parser.add_argument("--supporting-xlsx", type=Path)
    parser.add_argument("--expect-sheet", action="append", default=[])
    parser.add_argument("--json-out", type=Path)
    args = parser.parse_args()

    document = Document(args.docx)
    paragraphs = list(all_paragraphs(document))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    findings: list[dict[str, object]] = []

    if "{" in text or "}" in text:
        findings.append({"severity": "blocker", "check": "unresolved_brace_comment"})

    if args.expect_double_spacing:
        failures = [
            index
            for index, paragraph in enumerate(paragraphs)
            if effective_line_spacing_rule(paragraph) != WD_LINE_SPACING.DOUBLE
        ]
        if failures:
            findings.append({"severity": "blocker", "check": "double_spacing", "paragraph_indices": failures[:50], "failure_count": len(failures)})

    if args.expect_plain_black:
        run_failures = []
        run_index = 0
        for paragraph in paragraphs:
            for run in paragraph.runs:
                if not run.text:
                    continue
                bold = effective_run_font_value(run, paragraph, "bold")
                italic = effective_run_font_value(run, paragraph, "italic")
                underline = effective_run_font_value(run, paragraph, "underline")
                color_value = effective_run_color(run, paragraph)
                if bold is True or italic is True or underline not in {None, False}:
                    run_failures.append(run_index)
                elif color_value is not None and color_value != "000000":
                    run_failures.append(run_index)
                run_index += 1
        if run_failures:
            findings.append({"severity": "blocker", "check": "plain_black_text", "run_indices": run_failures[:50], "failure_count": len(run_failures)})

    table_shapes = [(len(table.rows), len(table.columns)) for table in document.tables]
    if args.expect_table_count is not None and len(document.tables) != args.expect_table_count:
        findings.append(
            {
                "severity": "blocker",
                "check": "table_count",
                "expected": args.expect_table_count,
                "actual": len(document.tables),
            }
        )
    if args.expect_table_shape:
        expected_shapes = []
        for value in args.expect_table_shape:
            match = re.fullmatch(r"(\d+)[xX](\d+)", value)
            if not match:
                raise SystemExit(f"Invalid --expect-table-shape value: {value!r}; expected ROWSxCOLS")
            expected_shapes.append((int(match.group(1)), int(match.group(2))))
        if expected_shapes != table_shapes:
            findings.append(
                {
                    "severity": "blocker",
                    "check": "table_shapes",
                    "expected": expected_shapes,
                    "actual": table_shapes,
                }
            )
    for table_index, table in enumerate(document.tables, start=1):
        repeated_header_rows = []
        for row_index, row in enumerate(table.rows[1:], start=2):
            first_cell = row.cells[0].text.strip().casefold()
            if first_cell in {"category", "category a", "fragment"}:
                repeated_header_rows.append(row_index)
        if repeated_header_rows:
            findings.append(
                {
                    "severity": "blocker",
                    "check": "possible_concatenated_table",
                    "table_index": table_index,
                    "repeated_header_rows": repeated_header_rows,
                }
            )

    workbook_sheets = []
    if args.supporting_xlsx:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise SystemExit("openpyxl is required for supporting-workbook auditing") from exc
        workbook = load_workbook(args.supporting_xlsx, read_only=True, data_only=True)
        workbook_sheets = workbook.sheetnames
        missing_sheets = [sheet for sheet in args.expect_sheet if sheet not in workbook_sheets]
        if missing_sheets:
            findings.append(
                {
                    "severity": "blocker",
                    "check": "supporting_workbook_sheets",
                    "missing": missing_sheets,
                    "actual": workbook_sheets,
                }
            )

    image_rows = []
    for index, shape in enumerate(document.inline_shapes, start=1):
        relationship_id = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
        blob = document.part.related_parts[relationship_id].blob
        with Image.open(BytesIO(blob)) as image:
            pixel_ratio = image.width / image.height
        display_ratio = shape.width / shape.height
        image_rows.append({"index": index, "pixel_ratio": pixel_ratio, "display_ratio": display_ratio, "ratio_error": abs(pixel_ratio - display_ratio)})
        if abs(pixel_ratio - display_ratio) >= 1e-4:
            findings.append({"severity": "blocker", "check": "image_aspect_ratio", "image_index": index, "pixel_ratio": pixel_ratio, "display_ratio": display_ratio})

    captions = [paragraph.text for paragraph in document.paragraphs if re.match(r"^Figure\s+\d+\.", paragraph.text)]
    for caption in captions:
        labels = re.findall(r"\(([A-Z])\)", caption)
        if labels and labels != [chr(ord("A") + offset) for offset in range(len(labels))]:
            findings.append({"severity": "warning", "check": "panel_caption_sequence", "caption": caption})

    result = {
        "document": str(args.docx),
        "paragraph_count": len(paragraphs),
        "table_count": len(document.tables),
        "table_shapes": table_shapes,
        "image_count": len(document.inline_shapes),
        "figure_caption_count": len(captions),
        "supporting_workbook_sheets": workbook_sheets,
        "images": image_rows,
        "findings": findings,
        "verdict": "PASS" if not any(item["severity"] == "blocker" for item in findings) else "BLOCK",
    }
    output = json.dumps(result, ensure_ascii=False, indent=2)
    if args.json_out:
        args.json_out.parent.mkdir(parents=True, exist_ok=True)
        args.json_out.write_text(output, encoding="utf-8")
    print(output)
    raise SystemExit(0 if result["verdict"] == "PASS" else 1)


if __name__ == "__main__":
    main()
